#!/usr/bin/env python3
"""Generate research/Literature_Survey_V3.docx from research/survey50.csv.

Landscape A4, title page, short introduction (survey/intro.md, with counts filled in from the
CSV), then one table grouped by category (category heading rows), header row repeated on every
page. All text uses one font family. Usage: python make_docx.py
"""
import csv
import datetime
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_survey import CATEGORIES

HERE = Path(__file__).resolve().parent
RES = HERE.parent
FONT = "Calibri"
HEADERS = ["S.No", "Title", "Authors", "Journal/Conference", "Vol/Issue/Pages/Year", "DOI",
           "Work done", "Advantages", "Disadvantages", "Relevance"]
WIDTHS_CM = [1.25, 3.3, 3.0, 2.7, 2.1, 2.5, 3.25, 2.9, 2.9, 2.8]   # sum 26.7 = A4 landscape - 2 x 1.5 cm


def set_font(style, size, bold=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def cell_text(cell, text, size=8, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = FONT
    run.bold = bold


def vip(r):
    return (f"Vol. {r['volume']}, No. {r['issue']}, pp. {r['pages']}, {r['year']}")


def main():
    rows = list(csv.DictReader((RES / "survey50.csv").open(encoding="utf-8")))
    doc = Document()

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.5))
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    set_font(doc.styles["Normal"], 11)
    for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 12)):
        set_font(doc.styles[name], size, bold=True)
        doc.styles[name].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    # ---- title page
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_paragraph("Literature Survey", style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line, size in (("A Timing-Predictable INT8 CNN Accelerator on the AMD Kria KV260 (V3)", 16),
                       (f"{len(rows)} verified papers in {len(CATEGORIES)} categories", 12),
                       (datetime.date.today().strftime("%d %B %Y"), 12)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.size = Pt(size)
    doc.paragraphs[-1].runs[-1].add_break(WD_BREAK.PAGE)

    # ---- introduction
    cat_count = Counter(r["category"] for r in rows)
    basis = Counter(r["basis"] for r in rows)
    fields = {
        "n": len(rows), "n_full": basis.get("full text", 0), "n_abs": basis.get("abstract", 0),
        "categories": "; ".join(f"({c}) {name}: {cat_count.get(c, 0)}"
                                for c, (name, _) in CATEGORIES.items()),
    }
    doc.add_heading("Introduction", level=1)
    for para in (HERE / "intro.md").read_text(encoding="utf-8").split("\n\n"):
        para = para.strip()
        if para and not para.startswith("#"):
            p = doc.add_paragraph(" ".join(para.split()).format(**fields))
            p.paragraph_format.space_after = Pt(6)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- table
    doc.add_heading("Survey table", level=1)
    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")           # repeat header row on every page
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for i, h in enumerate(HEADERS):
        cell_text(hdr.cells[i], h, size=8.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(hdr.cells[i], "1F2A44")
        hdr.cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    current = None
    for r in rows:
        if r["category"] != current:
            current = r["category"]
            cells = table.add_row().cells
            merged = cells[0].merge(cells[-1])
            cell_text(merged, f"Category {current}: {CATEGORIES[current][0]} "
                              f"({cat_count[current]} papers)", size=9.5, bold=True)
            shade(merged, "DCE3EF")
        cells = table.add_row().cells
        values = [r["s_no"], r["title"], r["authors"], r["venue"], vip(r), r["doi"],
                  r["work_done"], r["advantages"], r["disadvantages"], r["relevance"]]
        for i, v in enumerate(values):
            cell_text(cells[i], str(v), align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else None)
        basis_run = cells[6].paragraphs[0].add_run(f" [Basis: {r['basis']}]")
        basis_run.font.size, basis_run.font.name, basis_run.italic = Pt(7), FONT, True

    for i, w in enumerate(WIDTHS_CM):
        table.columns[i].width = Cm(w)
    for row in table.rows:
        if len({id(c._tc) for c in row.cells}) == 1:      # merged category heading row
            row.cells[0].width = Cm(sum(WIDTHS_CM))
            continue
        for i, w in enumerate(WIDTHS_CM):
            row.cells[i].width = Cm(w)

    out = RES / "Literature_Survey_V3.docx"
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
